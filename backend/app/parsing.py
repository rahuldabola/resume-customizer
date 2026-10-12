"""Turn an uploaded PDF or DOCX into plain text."""

import io
import re
from dataclasses import dataclass

import pdfplumber
from docx import Document


class ResumeFileError(ValueError):
    """Raised with a message that is safe to show to the user."""


@dataclass
class ExtractedResume:
    text: str
    pages: int


def detect_kind(filename: str, data: bytes) -> str:
    """Check both the extension and the file signature, so a renamed file is rejected."""
    name = filename.lower()
    if name.endswith(".pdf"):
        if not data.startswith(b"%PDF"):
            raise ResumeFileError("This file has a .pdf extension but is not a valid PDF.")
        return "pdf"
    if name.endswith(".docx"):
        if not data.startswith(b"PK"):
            raise ResumeFileError("This file has a .docx extension but is not a valid Word document.")
        return "docx"
    if name.endswith(".doc"):
        raise ResumeFileError("Old .doc files are not supported. Save the resume as .docx or PDF.")
    raise ResumeFileError("Unsupported file type. Upload a PDF or DOCX resume.")


def _extract_pdf(data: bytes) -> ExtractedResume:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [page.extract_text(x_tolerance=1.5, y_tolerance=3) or "" for page in pdf.pages]
    except Exception as exc:  # pdfminer raises many different types for broken files
        raise ResumeFileError("Could not read this PDF. It may be corrupted or password-protected.") from exc
    return ExtractedResume(text="\n\n".join(pages), pages=len(pages))


def _extract_docx(data: bytes) -> ExtractedResume:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeFileError("Could not read this Word document. It may be corrupted.") from exc

    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        lines.append(f"- {text}" if "list" in style else text)
    # Many resume templates lay out content in tables.
    for table in doc.tables:
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                value = cell.text.strip()
                if value and value not in cells:  # merged cells repeat their text
                    cells.append(value)
            if cells:
                lines.append(" | ".join(cells))
    return ExtractedResume(text="\n".join(lines), pages=1)


def clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_resume_text(filename: str, data: bytes) -> ExtractedResume:
    kind = detect_kind(filename, data)
    extracted = _extract_pdf(data) if kind == "pdf" else _extract_docx(data)
    extracted.text = clean_text(extracted.text)
    return extracted
