"""Render a structured resume as TXT, DOCX or PDF (single-column, ATS-friendly layout)."""

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from fpdf import FPDF

from .schemas import Resume

# DejaVu covers non-Latin-1 characters (smart quotes, en dashes, accented names). The Docker
# image installs it; elsewhere we fall back to a core font with characters transliterated.
_FONT_DIRS = [Path("/usr/share/fonts/truetype/dejavu"), Path(__file__).parent / "fonts"]


def _sections(resume: Resume) -> list[tuple[str, list[tuple[str, str, list[str]]]]]:
    """Common layout used by every format: (heading, [(title, meta, bullets)])."""
    out: list[tuple[str, list[tuple[str, str, list[str]]]]] = []
    if resume.summary.strip():
        out.append(("Summary", [("", "", [resume.summary.strip()])]))
    if resume.skills:
        out.append(("Skills", [(g.category, ", ".join(g.items), []) for g in resume.skills if g.items]))
    if resume.experience:
        out.append(
            (
                "Experience",
                [
                    (
                        " - ".join(x for x in (e.title, e.company) if x),
                        " | ".join(x for x in (e.location, e.dates) if x),
                        e.bullets,
                    )
                    for e in resume.experience
                ],
            )
        )
    if resume.projects:
        out.append(("Projects", [(p.name, p.subtitle, p.bullets) for p in resume.projects]))
    if resume.education:
        out.append(
            (
                "Education",
                [
                    (
                        " - ".join(x for x in (ed.degree, ed.institution) if x),
                        ed.dates,
                        [ed.details] if ed.details else [],
                    )
                    for ed in resume.education
                ],
            )
        )
    for extra in resume.extra_sections:
        if extra.items:
            out.append((extra.title, [("", "", extra.items)]))
    return out


def to_txt(resume: Resume) -> bytes:
    lines = [resume.name.upper()]
    if resume.headline:
        lines.append(resume.headline)
    if resume.contact:
        lines.append(" | ".join(resume.contact))
    for heading, entries in _sections(resume):
        lines += ["", heading.upper(), "-" * len(heading)]
        for title, meta, bullets in entries:
            if heading == "Skills":
                lines.append(f"{title}: {meta}")
                continue
            if title or meta:
                lines.append("  |  ".join(x for x in (title, meta) if x))
            prefix = "" if heading == "Summary" else "- "
            lines += [f"{prefix}{b}" for b in bullets]
    return ("\n".join(lines) + "\n").encode("utf-8")


def to_docx(resume: Resume) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(50)
        section.top_margin = section.bottom_margin = Pt(40)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(resume.name)
    run.bold = True
    run.font.size = Pt(20)
    for line in (resume.headline, " | ".join(resume.contact)):
        if line:
            p = doc.add_paragraph(line)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)

    for heading, entries in _sections(resume):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        hr = h.add_run(heading.upper())
        hr.bold = True
        hr.font.size = Pt(11.5)
        hr.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8A)
        for title, meta, bullets in entries:
            if heading == "Skills":
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                p.add_run(f"{title}: ").bold = True
                p.add_run(meta)
                continue
            if title or meta:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                p.add_run(title).bold = True
                if meta:
                    p.add_run(f"   {meta}").italic = True
            for b in bullets:
                bp = doc.add_paragraph(b, style=None if heading == "Summary" else "List Bullet")
                bp.paragraph_format.space_after = Pt(1)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_TRANSLITERATE = str.maketrans(
    {"–": "-", "—": "-", "‘": "'", "’": "'", "“": '"', "”": '"', "•": "-", "…": "...", " ": " ", "→": "->"}
)


class _PDF(FPDF):
    def __init__(self) -> None:
        super().__init__(format="A4")
        self.set_margins(16, 14, 16)
        self.set_auto_page_break(True, margin=14)
        self.unicode = False
        for folder in _FONT_DIRS:
            regular, bold = folder / "DejaVuSans.ttf", folder / "DejaVuSans-Bold.ttf"
            if regular.exists() and bold.exists():
                self.add_font("Body", "", str(regular))
                self.add_font("Body", "B", str(bold))
                self.add_font("Body", "I", str(regular))
                self.unicode = True
                break

    def font(self, style: str = "", size: float = 10) -> None:
        self.set_font("Body" if self.unicode else "Helvetica", style, size)

    def clean(self, text: str) -> str:
        if self.unicode:
            return text
        return text.translate(_TRANSLITERATE).encode("latin-1", "replace").decode("latin-1")


def to_pdf(resume: Resume) -> bytes:
    pdf = _PDF()
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin

    pdf.font("B", 18)
    pdf.cell(width, 9, pdf.clean(resume.name), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.font("", 10)
    for line in (resume.headline, " | ".join(resume.contact)):
        if line:
            pdf.multi_cell(width, 5, pdf.clean(line), align="C", new_x="LMARGIN", new_y="NEXT")

    for heading, entries in _sections(resume):
        pdf.ln(3)
        pdf.font("B", 11)
        pdf.set_text_color(31, 58, 138)
        pdf.cell(width, 6, pdf.clean(heading.upper()), new_x="LMARGIN", new_y="NEXT")
        pdf.set_draw_color(180, 190, 210)
        pdf.line(pdf.l_margin, pdf.get_y(), pdf.l_margin + width, pdf.get_y())
        pdf.ln(1.5)
        pdf.set_text_color(20, 20, 20)
        for title, meta, bullets in entries:
            if heading == "Skills":
                pdf.font("B", 10)
                label = pdf.clean(f"{title}: ")
                pdf.write(5, label)
                pdf.font("", 10)
                pdf.write(5, pdf.clean(meta))
                pdf.ln(5.5)
                continue
            if title or meta:
                pdf.font("B", 10.5)
                pdf.multi_cell(width, 5.5, pdf.clean(title), new_x="LMARGIN", new_y="NEXT")
                if meta:
                    pdf.font("I", 9.5)
                    pdf.set_text_color(90, 90, 90)
                    pdf.multi_cell(width, 4.8, pdf.clean(meta), new_x="LMARGIN", new_y="NEXT")
                    pdf.set_text_color(20, 20, 20)
            pdf.font("", 10)
            for b in bullets:
                if heading == "Summary":
                    pdf.multi_cell(width, 5, pdf.clean(b), new_x="LMARGIN", new_y="NEXT")
                    continue
                pdf.set_x(pdf.l_margin + 2)
                pdf.cell(4, 5, "-" if not pdf.unicode else "•")
                pdf.multi_cell(width - 6, 5, pdf.clean(b), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1.2)
    return bytes(pdf.output())


EXPORTERS = {
    "txt": (to_txt, "text/plain; charset=utf-8"),
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pdf": (to_pdf, "application/pdf"),
}
