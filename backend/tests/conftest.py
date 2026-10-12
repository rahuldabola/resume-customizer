import io
from pathlib import Path

import pytest
from docx import Document

from app.schemas import (
    EducationItem,
    ExperienceItem,
    ExtraSection,
    JDAnalysis,
    ProjectItem,
    Resume,
    SkillGroup,
)

SAMPLES = Path(__file__).resolve().parents[2] / "samples"


@pytest.fixture
def sample_pdf() -> bytes:
    return (SAMPLES / "sample_resume.pdf").read_bytes()


@pytest.fixture
def sample_jd() -> str:
    return (SAMPLES / "sample_job_description.txt").read_text(encoding="utf-8")


@pytest.fixture
def sample_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("Priya Sharma")
    doc.add_paragraph("priya@example.com | +91 90000 00000")
    doc.add_paragraph("Backend developer with four years of Python and Django experience. " * 3)
    doc.add_paragraph("Built payment APIs in Django REST Framework", style="List Bullet")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Python, Django, PostgreSQL"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def resume() -> Resume:
    return Resume(
        name="Aarav Mehta",
        headline="Software Engineer",
        contact=["aarav@example.com", "+91 98765 43210", "Noida"],
        summary="Software engineer building web apps.",
        skills=[SkillGroup(category="Skills", items=["Python", "React", "Postgres", "Docker"])],
        experience=[
            ExperienceItem(
                title="Software Engineer",
                company="Brightpath",
                location="Noida",
                dates="Jan 2024 - Present",
                bullets=[
                    "Made endpoints in Flask for the mobile team.",
                    "Added an OpenAI-based reply feature that cut handling time by 18%.",
                ],
            )
        ],
        projects=[ProjectItem(name="DocChat", subtitle="Python", bullets=["Built a RAG chatbot over PDFs."])],
        education=[EducationItem(degree="B.Tech CSE", institution="Galgotias", dates="2019 - 2023", details="")],
        extra_sections=[ExtraSection(title="Certifications", items=["AWS Cloud Practitioner"])],
    )


@pytest.fixture
def jd() -> JDAnalysis:
    return JDAnalysis(
        role_title="Full Stack Engineer",
        company="Lumora",
        seniority="Mid",
        must_have=["Python", "React / Next.js", "PostgreSQL", "Kubernetes"],
        nice_to_have=["Redis"],
        soft_skills=["ownership"],
        key_responsibilities=["build REST APIs"],
    )
